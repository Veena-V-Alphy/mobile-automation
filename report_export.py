import glob
import html
import json
import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

STATUS_FILL = {
    "passed": "C6EFCE",
    "failed": "FFC7CE",
    "broken": "FFEB9C",
    "skipped": "D9D9D9",
}

HEADERS = ["Suite", "Test Name", "Status", "Start Time", "Duration (s)", "Steps", "Failure Message"]


def _load_results(results_dir):
    rows = []
    for path in glob.glob(os.path.join(results_dir, "*-result.json")):
        with open(path, encoding="utf-8") as f:
            result = json.load(f)

        suite = next(
            (label["value"] for label in result.get("labels", []) if label["name"] == "suite"),
            "",
        )
        start_ms = result.get("start", 0)
        stop_ms = result.get("stop", start_ms)
        steps = "; ".join(
            f'{step["name"]} ({step["status"]})' for step in result.get("steps", [])
        )
        failure_message = (result.get("statusDetails") or {}).get("message", "")

        rows.append({
            "suite": suite,
            "name": result.get("name", ""),
            "status": result.get("status", ""),
            "start": datetime.fromtimestamp(start_ms / 1000) if start_ms else None,
            "duration": (stop_ms - start_ms) / 1000,
            "steps": steps,
            "failure_message": failure_message,
        })

    rows.sort(key=lambda r: r["start"] or datetime.min)
    return rows


def export_to_excel(results_dir, out_path):
    rows = _load_results(results_dir)

    wb = Workbook()
    ws = wb.active
    ws.title = "Allure Results"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row in rows:
        ws.append([
            row["suite"],
            row["name"],
            row["status"],
            row["start"].strftime("%Y-%m-%d %H:%M:%S") if row["start"] else "",
            round(row["duration"], 2),
            row["steps"],
            row["failure_message"],
        ])
        fill_color = STATUS_FILL.get(row["status"])
        if fill_color:
            ws.cell(row=ws.max_row, column=3).fill = PatternFill(
                start_color=fill_color, end_color=fill_color, fill_type="solid"
            )

    for column_cells in ws.columns:
        length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
        ws.column_dimensions[column_cells[0].column_letter].width = min(length + 2, 60)

    wb.save(out_path)
    return out_path


def export_to_html(results_dir, out_path):
    rows = _load_results(results_dir)

    table_rows = []
    for row in rows:
        status_color = {
            "passed": "#c6efce",
            "failed": "#ffc7ce",
            "broken": "#ffeb9c",
            "skipped": "#d9d9d9",
        }.get(row["status"], "#ffffff")
        table_rows.append(
            "<tr>"
            f"<td>{html.escape(row['suite'])}</td>"
            f"<td>{html.escape(row['name'])}</td>"
            f"<td style='background:{status_color}'>{html.escape(row['status'])}</td>"
            f"<td>{row['start'].strftime('%Y-%m-%d %H:%M:%S') if row['start'] else ''}</td>"
            f"<td>{round(row['duration'], 2)}</td>"
            f"<td>{html.escape(row['steps'])}</td>"
            f"<td>{html.escape(row['failure_message'])}</td>"
            "</tr>"
        )

    document = f"""<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<title>Allure Results Summary</title>
<style>
  body {{ font-family: Arial, sans-serif; margin: 24px; }}
  table {{ border-collapse: collapse; width: 100%; }}
  th, td {{ border: 1px solid #999; padding: 6px 10px; text-align: left; font-size: 13px; vertical-align: top; }}
  th {{ background: #333; color: #fff; }}
</style>
</head>
<body>
<h2>Allure Results Summary</h2>
<table>
<tr>{"".join(f"<th>{h}</th>" for h in HEADERS)}</tr>
{"".join(table_rows)}
</table>
</body>
</html>
"""
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(document)
    return out_path


if __name__ == "__main__":
    export_to_excel("allure-results", "allure-report-summary.xlsx")
    export_to_html("allure-results", "allure-report-summary.html")
